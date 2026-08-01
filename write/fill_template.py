#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""将修改后的章节内容填入期刊模板 — 简洁重建版"""
import sys, io, os, re
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')
from docx import Document
from docx.shared import Pt
from docx.oxml.ns import qn

BASE = r"C:\Users\86152\PycharmProjects\2 +ao arctic_seaice_prediction lstm SIE"
TEMPLATE = os.path.join(BASE, "write", "模板 基于LSTM的多源气候指数的北极海冰面积预测研究 (已自动恢复).docx")
OUTPUT = os.path.join(BASE, "write", "基于LSTM的多源气候指数的北极海冰面积预测研究_填入终稿_v5.docx")
CHAPTERS_DIR = os.path.join(BASE, "write", "chapters")

CHAPTER_FILES = [
    "01_数据与方法.md",
    "02_实验方案设计.md",
    "03_结果与分析.md",
]

def clean(s):
    return s.replace("**", "").strip()

def main():
    print("=" * 60)
    print("加载模板...")
    doc = Document(TEMPLATE)

    body = doc.element.body
    all_p = body.findall(qn("w:p"))

    # 查找章节边界
    ch2_idx = None  # "数据与方法" heading
    ch4_idx = None  # "结论与展望" heading

    for i, p in enumerate(all_p):
        text = "".join(p.itertext()).strip()
        if "数据与方法" in text and ch2_idx is None:
            ch2_idx = i
        if "结论与展望" in text:
            ch4_idx = i
            break

    if ch2_idx is None or ch4_idx is None:
        print(f"ERROR: ch2_idx={ch2_idx}, ch4_idx={ch4_idx}")
        for i, p in enumerate(all_p):
            text = "".join(p.itertext()).strip()
            if text:
                print(f"  [{i}] {text[:80]}")
        return

    print(f"第二章边界: 元素索引 {ch2_idx}")
    print(f"第四章边界: 元素索引 {ch4_idx}")

    # 删除 ch2_idx 到 ch4_idx-1 之间的所有元素
    for i in range(ch4_idx - 1, ch2_idx - 1, -1):
        body.remove(all_p[i])

    remaining = body.findall(qn("w:p"))
    print(f"删除后剩余段落元素: {len(remaining)}")

    # 找到插入点
    insert_after = all_p[ch2_idx - 1] if ch2_idx > 0 else None
    print(f"插入点: {insert_after is not None}")

    # 解析 markdown 文件
    print("\n解析修改后的章节文件...")
    content_items = []

    for fname in CHAPTER_FILES:
        fpath = os.path.join(CHAPTERS_DIR, fname)
        with open(fpath, "r", encoding="utf-8") as f:
            md = f.read()

        lines = md.split("\n")
        i = 0
        while i < len(lines):
            s = lines[i].strip()
            if not s:
                i += 1
                continue

            # 跳过第一个顶层标题 (# 1 数据与方法) — 已手动插入
            # 保留 # 2 实验方案设计 和 # 3 结果与分析 作为 Heading 1
            if s.startswith("# ") and not s.startswith("## "):
                h_text = clean(s.split(" ", 1)[1] if " " in s else s[2:])
                if h_text.startswith("1 ") or h_text.startswith("1."):
                    i += 1
                    continue
                else:
                    content_items.append(("h1", h_text))
                    i += 1
                    continue

            if s.startswith("## "):
                content_items.append(("h2", clean(s[3:])))
                i += 1
            elif s.startswith("### "):
                content_items.append(("h3", clean(s[4:])))
                i += 1
            elif s.startswith("$$"):
                i += 1
                continue
            elif s.startswith("|"):
                rows = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    rl = lines[i].strip()
                    if not re.match(r"^\|[\s\-:|]+\|$", rl):
                        cells = [c.strip() for c in rl.split("|")[1:-1]]
                        rows.append(cells)
                    i += 1
                if rows:
                    content_items.append(("table", rows))
            else:
                content_items.append(("text", clean(s)))
                i += 1

    print(f"共生成 {len(content_items)} 个内容段落")

    # 创建并插入新段落
    print("插入内容到模板...")
    from lxml import etree

    W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"

    def make_p(text, style_id, bold=False):
        p = etree.Element(qn("w:p"))
        pPr = etree.SubElement(p, qn("w:pPr"))
        pStyle = etree.SubElement(pPr, qn("w:pStyle"))
        pStyle.set(qn("w:val"), style_id)
        r = etree.SubElement(p, qn("w:r"))
        if bold:
            rPr2 = etree.SubElement(r, qn("w:rPr"))
            etree.SubElement(rPr2, qn("w:b"))
        t = etree.SubElement(r, qn("w:t"))
        t.text = text
        t.set(qn("xml:space"), "preserve")
        return p

    # 首先插入第二章标题
    ch2_heading = make_p("数据与方法", "1", bold=True)
    insert_after.addnext(ch2_heading)
    insert_after = ch2_heading

    count = 0
    for item_type, item_content in content_items:
        try:
            if item_type == "h1":
                new_p = make_p(item_content, "1", bold=True)
            elif item_type == "h2":
                new_p = make_p(item_content, "2", bold=True)
            elif item_type == "h3":
                new_p = make_p(item_content, "3", bold=True)
            elif item_type == "text":
                new_p = make_p(item_content, "Normal", bold=False)
            elif item_type == "table":
                # 创建真正的 docx 表格
                rows_data = item_content
                nrows = len(rows_data)
                ncols = len(rows_data[0]) if rows_data else 0
                if nrows == 0 or ncols == 0:
                    continue

                # 在文档中创建表格（会添加到文档末尾）
                tbl = doc.add_table(rows=nrows, cols=ncols, style="Table Grid")
                for ri, row_cells in enumerate(rows_data):
                    for ci, cell_text in enumerate(row_cells):
                        if ci < ncols:
                            cell = tbl.rows[ri].cells[ci]
                            cell.text = cell_text
                            for cp in cell.paragraphs:
                                for run in cp.runs:
                                    run.font.size = Pt(8)

                # 移动表格元素到正确位置
                tbl_elem = tbl._tbl
                # 从文档末尾移除
                body.remove(tbl_elem)
                # 插入到上一个元素之后
                insert_after.addnext(tbl_elem)
                # 更新上一个元素引用
                insert_after = tbl_elem
                continue
            else:
                continue

            insert_after.addnext(new_p)
            insert_after = new_p
            count += 1
            if count % 50 == 0:
                print(f"  已插入 {count} 个段落...")
        except Exception as e:
            print(f"  WARNING [{item_type}]: {e}")

    print(f"成功插入 {count} 个段落")

    # 保存
    print(f"\n保存到: {OUTPUT}")
    doc.save(OUTPUT)
    print("DONE!")

if __name__ == "__main__":
    main()
