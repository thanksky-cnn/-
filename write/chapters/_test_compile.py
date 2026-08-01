
import docx, sys, io, os
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8')

BASE = r'c:\Users\86152\PycharmProjects\2 +ao arctic_seaice_prediction lstm SIE'
PLOTS = os.path.join(BASE, 'outputs', 'plots', 'paper')
CHAPTERS = os.path.join(BASE, 'write', 'chapters')
TEMPLATE = os.path.join(BASE, 'write', '融合Cellpose与圆度判别的盐水冰微结构自动识别与参数提取方法.docx')
OUTPUT = os.path.join(BASE, 'write', '基于LSTM的多源气候指数的北极海冰面积预测研究_终稿.docx')

print('Creating final polished docx...')
doc = docx.Document(TEMPLATE)

# Clear template
body = doc.element.body
for p in body.findall(qn('w:p')):
    body.remove(p)

# Title page
doc.add_paragraph('基于LSTM的多源气候指数的北极海冰面积预测研究', style='Title')
p = doc.add_paragraph(style='作者')
p.add_run('谢天')
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p = doc.add_paragraph(style='单位')
p.add_run('（大连理工大学，海岸和近海工程国家重点实验室，辽宁  大连  116024）')
print('Title page OK')

# Save and check
doc.save(OUTPUT)
print(f'Saved test to: {OUTPUT}')
